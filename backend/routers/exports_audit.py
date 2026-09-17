"""Audit log + CSV/PDF Exports."""
from io import BytesIO
import re
from datetime import datetime, timedelta
from typing import Optional
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from deps import db, get_current_user
from helpers import csv_response, csv_response_pt, invoice_line_vab
from pdf_helpers import build_invoice_pdf, build_billing_orders_pdf, build_dashboard_pdf, build_annual_billing_pdf_grouped, build_proposal_follow_up_pdf
from routers.analytics import (
    by_commercial, by_client, by_manufacturer, _kpis_summary, _forecast_receiving,
    _forecast_invoicing, _vab_analysis, _active_orders,
)
from routers.finance import list_invoices

router = APIRouter()


def _xlsx_response(sheet_name: str, headers: list[str], rows: list[list], filename: str, widths: tuple[int, ...], money_columns: tuple[int, ...] = (), percent_columns: tuple[int, ...] = ()) -> Response:
    """Cria uma exportação Excel consistente para as listagens comerciais."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    worksheet.freeze_panes = "A2"
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)
    header_fill = PatternFill("solid", fgColor="002FA7")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for column in money_columns:
        for cell in worksheet.iter_cols(min_col=column, max_col=column, min_row=2):
            cell[0].number_format = '#,##0.00 [$€-pt-PT]'
    for column in percent_columns:
        for cell in worksheet.iter_cols(min_col=column, max_col=column, min_row=2):
            cell[0].number_format = "0%"
    for column, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(column)].width = width
    worksheet.auto_filter.ref = worksheet.dimensions
    buffer = BytesIO()
    workbook.save(buffer)
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/audit")
async def list_audit(entity: Optional[str] = None, entity_id: Optional[str] = None,
                     limit: int = 100, user: dict = Depends(get_current_user)):
    q = {}
    if entity:
        q["entity"] = entity
    if entity_id:
        q["entity_id"] = entity_id
    docs = await db.audit_log.find(q, {"_id": 0}).sort("at", -1).to_list(min(limit, 500))
    return {"rows": docs}


@router.get("/exports/leads.xlsx")
async def export_leads_xlsx(
    search: str = "",
    status: str = "",
    sort_by: str = "created_at",
    sort_dir: str = "desc",
    user: dict = Depends(get_current_user),
):
    """Exporta todas as leads correspondentes aos filtros da listagem."""
    query = {}
    if search.strip():
        pattern = {"$regex": re.escape(search.strip()), "$options": "i"}
        matching_clients = await db.clients.find({"name": pattern}, {"_id": 0, "id": 1}).to_list(5000)
        query["$or"] = [
            {field: pattern} for field in ("id", "description", "client_name_raw", "status")
        ] + [{"client_id": {"$in": [client["id"] for client in matching_clients]}}]
    statuses = [item for item in status.split(",") if item]
    if statuses:
        query["status"] = {"$in": statuses}

    leads = await db.leads.find(query, {"_id": 0}).to_list(10000)
    clients = {client["id"]: client.get("name", "") for client in await db.clients.find({}, {"_id": 0}).to_list(2000)}
    status_labels = {
        "nova": "Nova", "em_qualificacao": "Em qualificação",
        "convertida": "Convertida", "descartada": "Descartada",
    }

    def sort_value(lead: dict):
        if sort_by == "client":
            return (clients.get(lead.get("client_id")) or lead.get("client_name_raw") or "").lower()
        if sort_by == "value":
            return float(lead.get("estimated_value") or 0)
        if sort_by == "status":
            return status_labels.get(lead.get("status"), lead.get("status", "")).lower()
        return str(lead.get("created_at") or "")

    leads.sort(key=sort_value, reverse=sort_dir != "asc")
    return _xlsx_response(
        "Leads",
        ["Cliente", "Descrição", "Valor estimado (EUR)", "Estado", "Criada em", "Atualizada em"],
        [[
            clients.get(lead.get("client_id")) or lead.get("client_name_raw") or "-",
            lead.get("description") or "",
            float(lead.get("estimated_value") or 0),
            status_labels.get(lead.get("status"), lead.get("status", "")),
            str(lead.get("created_at") or "")[:19],
            str(lead.get("updated_at") or "")[:19],
        ] for lead in leads],
        "leads.xlsx",
        (34, 54, 22, 20, 22, 22),
        (3,),
    )


@router.get("/exports/opportunities.xlsx")
async def export_opportunities_xlsx(search: str = "", status: str = "", sort_by: str = "created_at", sort_dir: str = "desc", user: dict = Depends(get_current_user)):
    query = {}
    if search.strip():
        pattern = {"$regex": re.escape(search.strip()), "$options": "i"}
        matching_clients = await db.clients.find({"name": pattern}, {"_id": 0, "id": 1}).to_list(5000)
        query["$or"] = [{field: pattern} for field in ("id", "description", "status")] + [{"client_id": {"$in": [client["id"] for client in matching_clients]}}]
    statuses = [item for item in status.split(",") if item]
    if statuses:
        query["status"] = {"$in": statuses}
    opportunities = await db.opportunities.find(query, {"_id": 0}).to_list(10000)
    clients = {client["id"]: client.get("name", "") for client in await db.clients.find({}, {"_id": 0}).to_list(2000)}
    labels = {"aberta": "Aberta", "em_analise": "Em análise", "convertida": "Convertida", "perdida": "Perdida"}

    def sort_value(opportunity: dict):
        if sort_by == "client": return (clients.get(opportunity.get("client_id")) or "").lower()
        if sort_by == "value": return float(opportunity.get("estimated_value") or 0)
        if sort_by == "status": return labels.get(opportunity.get("status"), opportunity.get("status", "")).lower()
        return str(opportunity.get("created_at") or "")

    opportunities.sort(key=sort_value, reverse=sort_dir != "asc")
    return _xlsx_response("Oportunidades", ["Cliente", "Descrição", "Valor estimado (EUR)", "VAB estimado (EUR)", "Probabilidade", "Fecho previsto", "Prioridade", "Estado"], [[
        clients.get(opportunity.get("client_id")) or "-", opportunity.get("description") or "",
        float(opportunity.get("estimated_value") or 0), float(opportunity.get("estimated_vab") or 0),
        float(opportunity.get("probability") or 0) / 100, str(opportunity.get("expected_close_date") or "")[:10],
        opportunity.get("priority") or "", labels.get(opportunity.get("status"), opportunity.get("status", "")),
    ] for opportunity in opportunities], "oportunidades.xlsx", (34, 58, 24, 24, 16, 18, 16, 20), (3, 4), (5,))


@router.get("/exports/proposals.xlsx")
async def export_proposals_xlsx(search: str = "", status: str = "", sort_by: str = "created_at", sort_dir: str = "desc", user: dict = Depends(get_current_user)):
    query = {}
    if search.strip():
        pattern = {"$regex": re.escape(search.strip()), "$options": "i"}
        client_ids = [client["id"] for client in await db.clients.find({"name": pattern}, {"_id": 0, "id": 1}).to_list(5000)]
        opportunity_ids = [opportunity["id"] for opportunity in await db.opportunities.find({"description": pattern}, {"_id": 0, "id": 1}).to_list(5000)]
        query["$or"] = [{field: pattern} for field in ("id", "number", "description", "status")] + [{"client_id": {"$in": client_ids}}, {"opportunity_id": {"$in": opportunity_ids}}]
    statuses = [item for item in status.split(",") if item]
    if statuses:
        query["status"] = {"$in": statuses}
    proposals = await db.proposals.find(query, {"_id": 0}).to_list(10000)
    clients = {client["id"]: client.get("name", "") for client in await db.clients.find({}, {"_id": 0}).to_list(2000)}
    opportunities = {opportunity["id"]: opportunity.get("description", "") for opportunity in await db.opportunities.find({}, {"_id": 0}).to_list(10000)}
    labels = {"em_elaboracao": "Em elaboração", "enviada": "Enviada", "em_negociacao": "Em negociação", "ganha": "Ganha", "perdida": "Perdida", "expirada": "Expirada", "substituida": "Substituída"}

    def sort_value(proposal: dict):
        if sort_by == "number": return proposal.get("number") or ""
        if sort_by == "client": return (clients.get(proposal.get("client_id")) or "").lower()
        if sort_by == "opportunity": return (opportunities.get(proposal.get("opportunity_id")) or "").lower()
        if sort_by == "value": return float(proposal.get("total_net") or 0)
        if sort_by == "expected_close_date": return str(proposal.get("expected_close_date") or "")
        if sort_by == "status": return labels.get(proposal.get("status"), proposal.get("status", "")).lower()
        return str(proposal.get("created_at") or "")

    proposals.sort(key=sort_value, reverse=sort_dir != "asc")
    return _xlsx_response("Propostas", ["Número", "Versão", "Cliente", "Oportunidade", "Descrição", "Valor s/ IVA (EUR)", "VAB (EUR)", "Probabilidade", "Fecho previsto", "Estado"], [[
        proposal.get("number") or "", proposal.get("version") or 1, clients.get(proposal.get("client_id")) or "-",
        opportunities.get(proposal.get("opportunity_id")) or "", proposal.get("description") or "",
        float(proposal.get("total_net") or 0), float(proposal.get("total_vab") or 0), float(proposal.get("probability") if proposal.get("probability") is not None else 100) / 100,
        str(proposal.get("expected_close_date") or proposal.get("next_follow_up_date") or "")[:10], labels.get(proposal.get("status"), proposal.get("status", "")),
    ] for proposal in proposals], "propostas.xlsx", (21, 12, 34, 52, 52, 24, 20, 16, 18, 20), (6, 7), (8,))


@router.get("/exports/orders.xlsx")
async def export_orders_xlsx(search: str = "", status: str = "", sort_by: str = "order_date", sort_dir: str = "desc", user: dict = Depends(get_current_user)):
    query = {}
    if search.strip():
        pattern = {"$regex": re.escape(search.strip()), "$options": "i"}
        client_ids = [client["id"] for client in await db.clients.find({"name": pattern}, {"_id": 0, "id": 1}).to_list(5000)]
        query["$or"] = [{field: pattern} for field in ("id", "number", "po_number", "status")] + [{"client_id": {"$in": client_ids}}]
    statuses = [item for item in status.split(",") if item]
    if statuses:
        query["status"] = {"$in": statuses}
    orders = await db.orders.find(query, {"_id": 0}).to_list(10000)
    clients = {client["id"]: client.get("name", "") for client in await db.clients.find({}, {"_id": 0}).to_list(2000)}
    labels = {"aberta": "Aberta", "em_planeamento": "Em planeamento", "fulfilled": "Concluída", "cancelada": "Cancelada"}

    def sort_value(order: dict):
        if sort_by == "number": return order.get("number") or ""
        if sort_by == "client": return (clients.get(order.get("client_id")) or "").lower()
        if sort_by == "value": return float(order.get("total_net") or 0)
        if sort_by == "status": return labels.get(order.get("status"), order.get("status", "")).lower()
        return str(order.get("order_date") or order.get("created_at") or "")

    orders.sort(key=sort_value, reverse=sort_dir != "asc")
    return _xlsx_response("Encomendas", ["Número", "Cliente", "PO", "Valor s/ IVA (EUR)", "VAB (EUR)", "Estado", "Data"], [[
        order.get("number") or "", clients.get(order.get("client_id")) or "-", order.get("po_number") or "",
        float(order.get("total_net") or 0), float(order.get("total_vab") or 0), labels.get(order.get("status"), order.get("status", "")),
        str(order.get("order_date") or order.get("created_at") or "")[:10],
    ] for order in orders], "encomendas.xlsx", (21, 38, 24, 24, 20, 20, 16), (4, 5))


@router.get("/exports/invoices.xlsx")
async def export_invoices_xlsx(search: str = "", collection_status: str = "", sort_by: str = "issued_at", sort_dir: str = "desc", user: dict = Depends(get_current_user)):
    invoices = await list_invoices(user=user)
    labels = {"emitida": "Emitida", "parcialmente_recebida": "Parcialmente recebida", "recebida": "Recebida", "anulada": "Anulada"}
    collection_labels = {"em_atraso": "Em atraso", "por_receber": "Por receber", "recebida": "Recebida", "anulada": "Anulada"}
    term = search.strip().lower()
    selected_statuses = [item for item in collection_status.split(",") if item]
    invoices = [invoice for invoice in invoices if (
        not term or any(term in str(value or "").lower() for value in (invoice.get("order_number"), invoice.get("client_name"), invoice.get("external_invoice_number"), labels.get(invoice.get("status")), collection_labels.get(invoice.get("collection_status"))))
    ) and (not selected_statuses or invoice.get("collection_status") in selected_statuses)]
    invoices.sort(key=lambda invoice: invoice.get(sort_by) if isinstance(invoice.get(sort_by), (int, float)) else str(invoice.get(sort_by) or ""), reverse=sort_dir != "asc")
    return _xlsx_response("Faturas", ["Encomenda", "Cliente", "Data planeada", "Valor planeado (EUR)", "VAB planeado (EUR)", "Valor faturado (EUR)", "VAB faturado (EUR)", "N.º da fatura", "Estado", "Recebimento"], [[
        invoice.get("order_number") or "", invoice.get("client_name") or "", str(invoice.get("planned_date") or "")[:10],
        float(invoice.get("planned_value") or 0), float(invoice.get("planned_vab") or 0), float(invoice.get("total_net") or 0), float(invoice.get("billed_vab") or 0),
        invoice.get("external_invoice_number") or "", labels.get(invoice.get("status"), invoice.get("status", "")), collection_labels.get(invoice.get("collection_status"), ""),
    ] for invoice in invoices], "faturas.xlsx", (22, 38, 18, 25, 23, 25, 23, 22, 24, 20), (4, 5, 6, 7))


@router.get("/exports/invoices.csv")
async def export_invoices(user: dict = Depends(get_current_user)):
    active_order_ids = [o["id"] for o in await _active_orders()]
    invoices = await db.invoices.find(
        {"status": {"$ne": "anulada"}, "order_id": {"$in": active_order_ids}},
        {"_id": 0},
    ).sort("issued_at", -1).to_list(5000)
    clients = {c["id"]: c["name"] for c in await db.clients.find({}, {"_id": 0}).to_list(2000)}
    orders = {o["id"]: o["number"] for o in await db.orders.find({}, {"_id": 0}).to_list(2000)}
    rows = [{
        "numero_fatura_interna": i["number"],
        "numero_fatura_externa": i.get("external_invoice_number", ""),
        "data_emissao": i["issued_at"][:10],
        "cliente": clients.get(i["client_id"], ""),
        "encomenda": orders.get(i["order_id"], ""),
        "total_sem_iva": i["total_net"],
        "iva": i["total_vat"],
        "total_com_iva": i["total_gross"],
        "recebido": i.get("received_amount", 0),
        "saldo_em_aberto": max(0, float(i.get("total_gross") or 0) - float(i.get("received_amount") or 0)),
        "estado": i["status"],
    } for i in invoices]
    fields = ["numero_fatura_interna", "numero_fatura_externa", "data_emissao", "cliente", "encomenda", "total_sem_iva", "iva", "total_com_iva", "recebido", "saldo_em_aberto", "estado"]
    return csv_response_pt(
        rows,
        fields,
        "faturas.csv",
        money_fields=("total_sem_iva", "iva", "total_com_iva", "recebido", "saldo_em_aberto"),
        date_fields=("data_emissao",),
    )


@router.get("/exports/orders.csv")
async def export_orders(user: dict = Depends(get_current_user)):
    orders = sorted(await _active_orders(), key=lambda order: order.get("created_at", ""), reverse=True)
    clients = {c["id"]: c.get("name", "") for c in await db.clients.find({}, {"_id": 0}).to_list(2000)}
    proposals = {p["id"]: p for p in await db.proposals.find({}, {"_id": 0}).to_list(5000)}
    rows = []
    for order in orders:
        proposal = proposals.get(order.get("proposal_id"), {})
        rows.append({
            "numero_encomenda": order.get("number", ""),
            "cliente": clients.get(order.get("client_id"), ""),
            "valor_total_sem_iva": order.get("total_net", 0),
            "data_conversao_proposta": str(order.get("created_at") or "")[:19],
            "numero_proposta": proposal.get("number", ""),
        })
    fields = ["numero_encomenda", "cliente", "valor_total_sem_iva", "data_conversao_proposta", "numero_proposta"]
    return csv_response(rows, fields, "encomendas-com-propostas.csv")


@router.get("/exports/proposals.csv")
async def export_proposals(user: dict = Depends(get_current_user)):
    """Exporta todas as propostas com o respetivo contexto comercial."""
    def single_line(value) -> str:
        return " ".join(str(value or "").replace("\r", " ").replace("\n", " ").split())

    proposals = await db.proposals.find({}, {"_id": 0}).sort("created_at", -1).to_list(10000)
    clients = {c["id"]: c.get("name", "") for c in await db.clients.find({}, {"_id": 0}).to_list(2000)}
    opportunities = {o["id"]: o for o in await db.opportunities.find({}, {"_id": 0}).to_list(5000)}
    orders = {o["id"]: o for o in await db.orders.find({}, {"_id": 0}).to_list(10000)}

    rows = []
    for proposal in proposals:
        opportunity = opportunities.get(proposal.get("opportunity_id"), {})
        order = orders.get(proposal.get("converted_order_id"), {})
        rows.append({
            "numero_proposta": proposal.get("number", ""),
            "versao": proposal.get("version", ""),
            "data_criacao": str(proposal.get("created_at") or "")[:19],
            "cliente": clients.get(proposal.get("client_id"), ""),
            "descricao_oportunidade": single_line(opportunity.get("description", "")),
            "notas": single_line(proposal.get("notes", "")),
            "estado": proposal.get("status", ""),
            "valor_sem_iva": proposal.get("total_net", 0),
            "iva": proposal.get("total_vat", 0),
            "valor_com_iva": proposal.get("total_gross", 0),
            "vab": proposal.get("total_vab", 0),
            "data_conversao_encomenda": str(order.get("created_at") or "")[:19],
            "numero_encomenda": order.get("number", ""),
        })

    fields = [
        "numero_proposta", "versao", "data_criacao", "cliente", "descricao_oportunidade", "notas",
        "estado", "valor_sem_iva", "iva", "valor_com_iva", "vab",
        "data_conversao_encomenda", "numero_encomenda",
    ]
    return csv_response(rows, fields, "propostas.csv")


@router.get("/exports/timesheet.csv")
async def export_timesheet(project_id: Optional[str] = None, user: dict = Depends(get_current_user)):
    q = {"project_id": project_id} if project_id else {}
    entries = await db.time_entries.find(q, {"_id": 0}).sort("date", -1).to_list(10000)
    projects = {p["id"]: p["name"] for p in await db.projects.find({}, {"_id": 0}).to_list(1000)}
    rows = [{
        "data": e["date"], "projeto": projects.get(e["project_id"], ""),
        "developer": e["user_name"], "horas": e["hours"], "custo": e["cost"],
        "faturavel": "sim" if e["billable"] else "nao",
        "descricao": e.get("description", ""),
    } for e in entries]
    return csv_response(rows, ["data", "projeto", "developer", "horas", "custo", "faturavel", "descricao"], "timesheet.csv")


@router.get("/exports/reporting-commercial.csv")
async def export_reporting_commercial(year: int = Query(datetime.now().year, ge=2000, le=2100), user: dict = Depends(get_current_user)):
    data = await by_commercial(user, year)
    return csv_response(
        data["rows"],
        ["name", "role", "leads", "opps", "props", "won", "lost", "conversion_rate", "won_value", "won_vab", "orders_value", "orders_vab"],
        "reporting-comerciais.csv",
    )


@router.get("/exports/billing-orders.csv")
async def export_billing_orders(month: str, user: dict = Depends(get_current_user)):
    """Ordem de faturação: lista faturas emitidas no mês indicado (formato YYYY-MM).

    Uma linha por linha de fatura (mais granular que a export global).
    """
    if not month or len(month) != 7 or month[4] != "-":
        raise HTTPException(400, "Parâmetro 'month' deve ter o formato YYYY-MM")
    active_order_ids = [o["id"] for o in await _active_orders()]
    invoices = await db.invoices.find(
        {"issued_at": {"$regex": f"^{month}"}, "status": {"$ne": "anulada"}, "order_id": {"$in": active_order_ids}},
        {"_id": 0},
    ).sort("issued_at", 1).to_list(5000)
    clients = {c["id"]: c["name"] for c in await db.clients.find({}, {"_id": 0}).to_list(2000)}
    orders = {o["id"]: o for o in await db.orders.find({}, {"_id": 0}).to_list(2000)}
    plan_lines = {pl["id"]: pl for pl in await db.plan_lines.find({}, {"_id": 0}).to_list(5000)}

    rows = []
    for inv in invoices:
        order = orders.get(inv["order_id"], {})
        for il in inv.get("lines", []):
            pl = plan_lines.get(il.get("plan_line_id"), {})
            rows.append({
                "fatura": inv["number"],
                "data_emissao": inv["issued_at"][:10],
                "cliente": clients.get(inv["client_id"], ""),
                "encomenda": order.get("number", ""),
                "linha_descricao": il.get("description") or pl.get("description") or "",
                "linha_tipo": pl.get("type", ""),
                "valor_sem_iva": il.get("amount", 0),
                "iva_pct": inv.get("vat_pct", 0),
                "iva": round(il.get("amount", 0) * inv.get("vat_pct", 0) / 100, 2),
                "valor_com_iva": round(il.get("amount", 0) * (1 + inv.get("vat_pct", 0) / 100), 2),
                "estado_fatura": inv["status"],
                "recebido": inv.get("received_amount", 0),
            })
        if not inv.get("lines"):
            rows.append({
                "fatura": inv["number"],
                "data_emissao": inv["issued_at"][:10],
                "cliente": clients.get(inv["client_id"], ""),
                "encomenda": order.get("number", ""),
                "linha_descricao": "", "linha_tipo": "",
                "valor_sem_iva": inv["total_net"], "iva_pct": inv.get("vat_pct", 0),
                "iva": inv["total_vat"], "valor_com_iva": inv["total_gross"],
                "estado_fatura": inv["status"], "recebido": inv.get("received_amount", 0),
            })
    return csv_response(
        rows,
        ["fatura", "data_emissao", "cliente", "encomenda", "linha_descricao", "linha_tipo",
         "valor_sem_iva", "iva_pct", "iva", "valor_com_iva", "estado_fatura", "recebido"],
        f"ordem-faturacao-{month}.csv",
    )


async def _annual_billing_report(year: int, start_month: Optional[str] = None, end_month: Optional[str] = None):
    start_month = start_month or f"{year}-01"
    end_month = end_month or f"{year}-12"
    if not all(len(value) == 7 and value[4] == "-" for value in (start_month, end_month)):
        raise HTTPException(400, "Os meses devem usar o formato AAAA-MM")
    if start_month[:4] != str(year) or end_month[:4] != str(year) or start_month > end_month:
        raise HTTPException(400, "O periodo deve estar dentro do ano selecionado")
    active_orders = await _active_orders()
    active_order_ids = [order["id"] for order in active_orders]
    orders = {order["id"]: order for order in active_orders}
    clients = {client["id"]: client.get("name", "") for client in await db.clients.find({}, {"_id": 0}).to_list(2000)}
    plan_lines = await db.plan_lines.find(
        {"status": {"$ne": "cancelada"}, "order_id": {"$in": active_order_ids}},
        {"_id": 0},
    ).to_list(10000)
    plan_lines_by_id = {line.get("id"): line for line in plan_lines if line.get("id")}
    invoices = await db.invoices.find(
        {"status": {"$ne": "anulada"}, "order_id": {"$in": active_order_ids}},
        {"_id": 0},
    ).to_list(10000)
    invoices = [
        invoice for invoice in invoices
        if start_month <= str(invoice.get("issued_at") or invoice.get("created_at") or "")[:7] <= end_month
    ]
    invoice_by_plan = {}
    invoice_only_rows = []
    for invoice in invoices:
        invoice_number = invoice.get("number", "")
        lines = invoice.get("lines") or [{"amount": invoice.get("total_net", 0), "description": ""}]
        for line in lines:
            plan_id = line.get("plan_line_id")
            amount = float(line.get("amount") or 0)
            description = str(line.get("description") or "").strip()
            if plan_id:
                invoice_by_plan.setdefault(plan_id, []).append({
                    "number": invoice_number, "amount": amount, "description": description,
                    "vab_amount": line.get("vab_amount"),
                    "month": str(invoice.get("issued_at") or invoice.get("created_at") or "")[:7],
                    "order_id": invoice.get("order_id"), "client_id": invoice.get("client_id"),
                })
            else:
                order = orders.get(invoice.get("order_id"), {})
                total_net = float(order.get("total_net") or 0)
                vab = invoice_line_vab(line, order, amount)
                invoice_only_rows.append({
                    "mes": str(invoice.get("issued_at") or invoice.get("created_at") or "")[:7],
                    "encomenda": order.get("number", ""), "cliente": clients.get(invoice.get("client_id"), ""),
                    "item_fatura": f"{description or 'Fatura'} · {invoice_number}",
                    "planeado": 0, "faturado": amount, "por_faturar": 0,
                    "vab_faturado": vab, "vab_por_faturar": 0,
                })

    rows = []
    included_plan_ids = set()
    for plan_line in plan_lines:
        expected = str(plan_line.get("expected_date") or "")
        if not start_month <= expected[:7] <= end_month:
            continue
        included_plan_ids.add(plan_line.get("id"))
        order = orders.get(plan_line.get("order_id"), {})
        planned = float(plan_line.get("value") or 0)
        linked = invoice_by_plan.get(plan_line.get("id"), [])
        billed = sum(item["amount"] for item in linked)
        total_net = float(order.get("total_net") or 0)
        vab = float(order.get("total_vab") or 0) * planned / total_net if total_net else 0
        vab_billed = sum(
            invoice_line_vab(item, order, item.get("amount"))
            for item in linked
        )
        if not linked:
            vab_billed = 0.0
        invoice_label = ", ".join(
            f"{item['number']} · {item['description']}" if item['description'] else item['number']
            for item in linked
        )
        rows.append({
            "mes": expected[:7], "encomenda": order.get("number", ""),
            "cliente": clients.get(order.get("client_id"), ""),
            "item_fatura": f"{plan_line.get('description') or plan_line.get('type') or 'Sem descrição'}{f' · {invoice_label}' if invoice_label else ''}",
            "planeado": planned, "faturado": billed, "por_faturar": max(0, planned - billed),
            "vab_faturado": vab_billed, "vab_por_faturar": max(0, vab - vab_billed),
        })

    # Faturas do ano podem referenciar linhas cujo mês planeado pertence a
    # outro ano. Mantemos essas linhas no mês de emissão para alinhar o
    # total faturado com o Dashboard.
    for plan_id, linked in invoice_by_plan.items():
        if plan_id in included_plan_ids:
            continue
        plan_line = plan_lines_by_id.get(plan_id, {})
        for invoice_line in linked:
            order = orders.get(invoice_line.get("order_id") or plan_line.get("order_id"), {})
            amount = float(invoice_line.get("amount") or 0)
            total_net = float(order.get("total_net") or 0)
            vab_billed = invoice_line_vab(invoice_line, order, amount)
            item_name = plan_line.get("description") or plan_line.get("type") or "Sem descricao"
            description = invoice_line.get("description")
            if description and description != item_name:
                item_name = f"{item_name} - {description}"
            rows.append({
                "mes": invoice_line.get("month", ""),
                "encomenda": order.get("number", ""),
                "cliente": clients.get(invoice_line.get("client_id") or order.get("client_id"), ""),
                "item_fatura": f"{item_name} · {invoice_line.get('number', '')}",
                "planeado": 0, "faturado": amount, "por_faturar": 0,
                "vab_faturado": vab_billed, "vab_por_faturar": 0,
            })
    rows.extend(invoice_only_rows)
    summary = []
    first_month = int(start_month[5:7])
    last_month = int(end_month[5:7])
    for month in range(first_month, last_month + 1):
        key = f"{year}-{month:02d}"
        month_rows = [row for row in rows if row["mes"] == key]
        summary.append({
            "month": key,
            "planned_value": sum(row["planeado"] for row in month_rows),
            "billed_value": sum(row["faturado"] for row in month_rows),
            "remaining_value": sum(row["por_faturar"] for row in month_rows),
            "vab_billed_value": sum(row["vab_faturado"] for row in month_rows),
            "vab_remaining_value": sum(row["vab_por_faturar"] for row in month_rows),
        })
    return summary, rows


@router.get("/exports/billing-annual.csv")
async def export_billing_annual_csv(start_month: Optional[str] = None, end_month: Optional[str] = None, year: int = Query(datetime.now().year, ge=2000, le=2100), user: dict = Depends(get_current_user)):
    _summary, rows = await _annual_billing_report(year, start_month, end_month)
    fields = ["mes", "encomenda", "cliente", "item_fatura", "planeado", "faturado", "vab_faturado", "por_faturar", "vab_por_faturar"]
    return csv_response_pt(
        rows,
        fields,
        f"faturacao-anual-{year}.csv",
        money_fields=("planeado", "faturado", "vab_faturado", "por_faturar", "vab_por_faturar"),
        date_fields=("mes",),
    )


@router.get("/exports/billing-annual.pdf")
async def export_billing_annual_pdf(start_month: Optional[str] = None, end_month: Optional[str] = None, year: int = Query(datetime.now().year, ge=2000, le=2100), user: dict = Depends(get_current_user)):
    summary, rows = await _annual_billing_report(year, start_month, end_month)
    return build_annual_billing_pdf_grouped(year, summary, rows)


@router.get("/exports/billing-competence.csv")
async def export_billing_competence(month: Optional[str] = None, user: dict = Depends(get_current_user)):
    """Exporta a competência provável de cada linha faturada através do plano.

    Não altera datas. A competência é derivada de plan_lines.expected_date e
    fica separada da data real de emissão da fatura.
    """
    if month and (len(month) != 7 or month[4] != "-"):
        raise HTTPException(400, "Parâmetro 'month' deve ter o formato YYYY-MM")

    active_order_ids = [o["id"] for o in await _active_orders()]
    invoices = await db.invoices.find(
        {"status": {"$ne": "anulada"}, "order_id": {"$in": active_order_ids}},
        {"_id": 0},
    ).sort("issued_at", 1).to_list(5000)
    clients = {c["id"]: c.get("name", "") for c in await db.clients.find({}, {"_id": 0}).to_list(2000)}
    orders = {o["id"]: o for o in await db.orders.find({}, {"_id": 0}).to_list(2000)}
    plan_lines = {pl["id"]: pl for pl in await db.plan_lines.find({}, {"_id": 0}).to_list(5000)}
    invoice_ids = [invoice.get("id") for invoice in invoices if invoice.get("id")]
    payments = await db.payments.find(
        {"invoice_id": {"$in": invoice_ids}, "status": {"$ne": "anulado"}},
        {"_id": 0, "invoice_id": 1, "amount": 1},
    ).to_list(5000) if invoice_ids else []
    received_by_invoice = {}
    for payment in payments:
        invoice_id = payment.get("invoice_id")
        received_by_invoice[invoice_id] = received_by_invoice.get(invoice_id, 0) + float(payment.get("amount") or 0)

    rows = []
    for invoice in invoices:
        vat_pct = float(invoice.get("vat_pct") or 0)
        order = orders.get(invoice.get("order_id"), {})
        invoice_lines = invoice.get("lines") or []
        source_lines = invoice_lines or [{"amount": invoice.get("total_net", 0), "description": ""}]
        for invoice_line in source_lines:
            plan_line = plan_lines.get(invoice_line.get("plan_line_id"), {})
            expected_date = str(plan_line.get("expected_date") or "")[:10]
            competence = expected_date[:7] if expected_date else ""
            if month and competence != month:
                continue
            amount = float(invoice_line.get("amount") or 0)
            rows.append({
                "competencia": competence,
                "data_prevista_plano": expected_date,
                "fatura": invoice.get("number", ""),
                "data_emissao": str(invoice.get("issued_at") or invoice.get("created_at") or "")[:10],
                "cliente": clients.get(invoice.get("client_id"), ""),
                "encomenda": order.get("number", ""),
                "linha_descricao": invoice_line.get("description") or plan_line.get("description", ""),
                "linha_tipo": plan_line.get("type", ""),
                "valor_sem_iva": round(amount, 2),
                "iva_pct": vat_pct,
                "iva": round(amount * vat_pct / 100, 2),
                "valor_com_iva": round(amount * (1 + vat_pct / 100), 2),
                "recebido_fatura": round(received_by_invoice.get(invoice.get("id"), 0), 2),
                "estado_fatura": invoice.get("status", ""),
                "competencia_origem": "plan_lines.expected_date" if expected_date else "sem_data_no_plano",
            })

    columns = [
        "competencia", "data_prevista_plano", "fatura", "data_emissao", "cliente", "encomenda",
        "linha_descricao", "linha_tipo", "valor_sem_iva", "iva_pct", "iva", "valor_com_iva",
        "recebido_fatura", "estado_fatura", "competencia_origem",
    ]
    suffix = month or "todas"
    return csv_response(rows, columns, f"competencias-faturacao-{suffix}.csv")


# =========================================================
# PDF exports
# =========================================================
@router.get("/invoices/{iid}/pdf")
async def export_invoice_pdf(iid: str, user: dict = Depends(get_current_user)):
    inv = await db.invoices.find_one({"id": iid}, {"_id": 0})
    if not inv:
        raise HTTPException(404, "Fatura não encontrada")
    client = await db.clients.find_one({"id": inv["client_id"]}, {"_id": 0}) or {}
    order = await db.orders.find_one({"id": inv["order_id"]}, {"_id": 0}) or {}
    plan_lines_map = {pl["id"]: pl for pl in await db.plan_lines.find({"order_id": inv["order_id"]}, {"_id": 0}).to_list(1000)}
    payments = await db.payments.find({"invoice_id": iid}, {"_id": 0}).sort("paid_at", 1).to_list(500)
    return build_invoice_pdf(inv, client, order, plan_lines_map, payments)


@router.get("/exports/billing-orders.pdf")
async def export_billing_orders_pdf(month: str, user: dict = Depends(get_current_user)):
    if not month or len(month) != 7 or month[4] != "-":
        raise HTTPException(400, "Parâmetro 'month' deve ter o formato YYYY-MM")
    active_order_ids = [o["id"] for o in await _active_orders()]
    invoices = await db.invoices.find(
        {"issued_at": {"$regex": f"^{month}"}, "status": {"$ne": "anulada"}, "order_id": {"$in": active_order_ids}},
        {"_id": 0},
    ).sort("issued_at", 1).to_list(5000)
    invoice_ids = [invoice["id"] for invoice in invoices]
    payments = await db.payments.find(
        {"invoice_id": {"$in": invoice_ids}, "status": {"$ne": "anulado"}},
        {"_id": 0, "invoice_id": 1, "amount": 1},
    ).to_list(5000) if invoice_ids else []
    received_by_invoice = {}
    for payment in payments:
        received_by_invoice[payment["invoice_id"]] = received_by_invoice.get(payment["invoice_id"], 0) + float(payment.get("amount") or 0)
    for invoice in invoices:
        invoice["received_amount"] = round(received_by_invoice.get(invoice["id"], 0), 2)
    clients_map = {c["id"]: c["name"] for c in await db.clients.find({}, {"_id": 0}).to_list(2000)}
    orders_map = {o["id"]: o for o in await db.orders.find({}, {"_id": 0}).to_list(2000)}
    plan_lines_map = {pl["id"]: pl for pl in await db.plan_lines.find({}, {"_id": 0}).to_list(5000)}
    return build_billing_orders_pdf(month, invoices, clients_map, orders_map, plan_lines_map)


@router.get("/exports/proposals-follow-up.csv")
async def export_proposals_follow_up_csv(year: int = Query(datetime.now().year, ge=2000, le=2100), scope: str = Query("year", pattern="^(30d|quarter|year)$"), user: dict = Depends(get_current_user)):
    """Exportação da previsão de fecho: propostas ativas e oportunidades ponderadas."""
    proposals = await db.proposals.find(
        {"status": {"$in": ["em_elaboracao", "enviada", "em_negociacao"]}, "$or": [{"expected_close_date": {"$exists": True, "$ne": None}}, {"next_follow_up_date": {"$exists": True, "$ne": None}}]},
        {"_id": 0},
    ).to_list(10000)
    open_opportunities = await db.opportunities.find(
        {"status": {"$in": ["aberta", "em_analise"]}, "expected_close_date": {"$exists": True, "$ne": None}},
        {"_id": 0},
    ).to_list(10000)
    clients = {c["id"]: c.get("name", "") for c in await db.clients.find({}, {"_id": 0}).to_list(2000)}
    opportunities = {o["id"]: o.get("description", "") for o in await db.opportunities.find({}, {"_id": 0}).to_list(10000)}
    today = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
    anchor = today if year == today.year else datetime(year, 1, 1)
    thirty_days_end = anchor + timedelta(days=30)
    next_quarter_month = ((anchor.month - 1) // 3 + 1) * 3 + 1
    quarter_end = (datetime(year + 1, 1, 1) if next_quarter_month > 12 else datetime(year, next_quarter_month, 1)) - timedelta(days=1)
    year_end = datetime(year, 12, 31)
    status_labels = {
        "em_elaboracao": "Em elaboração", "enviada": "Enviada", "em_negociacao": "Em negociação",
        "ganha": "Ganha", "perdida": "Perdida", "expirada": "Expirada", "substituida": "Substituída",
    }
    opportunity_status_labels = {"aberta": "Aberta", "em_analise": "Em análise"}
    rows = []
    for proposal in proposals:
        follow_up = str(proposal.get("expected_close_date") or proposal.get("next_follow_up_date") or "")[:10]
        try:
            follow_up_date = datetime.strptime(follow_up, "%Y-%m-%d")
        except ValueError:
            continue
        if follow_up_date < anchor or follow_up_date > year_end:
            continue
        horizon = "Próximos 30 dias" if follow_up_date <= min(thirty_days_end, year_end) else ("Trimestre atual" if follow_up_date <= quarter_end else "Até final do ano")
        probability = min(100, max(0, float(proposal.get("probability") if proposal.get("probability") is not None else 100)))
        factor = probability / 100
        rows.append({
            "horizonte": horizon,
            "tipo": "Proposta",
            "registo": proposal.get("number", "-"),
            "cliente": clients.get(proposal.get("client_id"), "-"),
            "descricao": proposal.get("description") or opportunities.get(proposal.get("opportunity_id"), "-"),
            "valor_sem_iva": float(proposal.get("total_net") or 0) * factor,
            "vab": float(proposal.get("total_vab") or 0) * factor,
            "probabilidade": probability,
            "estado": status_labels.get(proposal.get("status"), proposal.get("status", "-")),
            "data_prevista_fecho": follow_up,
            "_date": follow_up_date,
        })
    for opportunity in open_opportunities:
        follow_up = str(opportunity.get("expected_close_date") or "")[:10]
        try:
            follow_up_date = datetime.strptime(follow_up, "%Y-%m-%d")
        except ValueError:
            continue
        if follow_up_date < anchor or follow_up_date > year_end:
            continue
        probability = min(100, max(0, float(opportunity.get("probability") or 0)))
        factor = probability / 100
        horizon = "Próximos 30 dias" if follow_up_date <= min(thirty_days_end, year_end) else ("Trimestre atual" if follow_up_date <= quarter_end else "Até final do ano")
        rows.append({
            "horizonte": horizon,
            "tipo": "Oportunidade",
            "registo": f"OPP-{str(opportunity.get('id') or '-')[:8].upper()}",
            "cliente": clients.get(opportunity.get("client_id"), "-"),
            "descricao": opportunity.get("description", "-"),
            "valor_sem_iva": float(opportunity.get("estimated_value") or 0) * factor,
            "vab": float(opportunity.get("estimated_vab") or 0) * factor,
            "probabilidade": probability,
            "estado": opportunity_status_labels.get(opportunity.get("status"), opportunity.get("status", "-")),
            "data_prevista_fecho": follow_up,
            "_date": follow_up_date,
        })
    allowed = {"30d": {"Próximos 30 dias"}, "quarter": {"Próximos 30 dias", "Trimestre atual"}, "year": {"Próximos 30 dias", "Trimestre atual", "Até final do ano"}}[scope]
    rows = [row for row in rows if row["horizonte"] in allowed]
    rows.sort(key=lambda row: row["_date"])
    for row in rows:
        row.pop("_date", None)
    return csv_response_pt(
        rows,
        ["horizonte", "tipo", "registo", "cliente", "descricao", "valor_sem_iva", "vab", "probabilidade", "estado", "data_prevista_fecho"],
        f"previsao-fecho-{year}-{scope}.csv",
        money_fields=("valor_sem_iva", "vab"),
        date_fields=("data_prevista_fecho",),
    )


@router.get("/exports/proposals-follow-up.xlsx")
async def export_proposals_follow_up_xlsx(year: int = Query(datetime.now().year, ge=2000, le=2100), scope: str = Query("year", pattern="^(30d|quarter|year)$"), user: dict = Depends(get_current_user)):
    """Exporta a previsão de fecho para Excel, mantendo valores ponderados."""
    proposals = await db.proposals.find(
        {"status": {"$in": ["em_elaboracao", "enviada", "em_negociacao"]}, "$or": [{"expected_close_date": {"$exists": True, "$ne": None}}, {"next_follow_up_date": {"$exists": True, "$ne": None}}]},
        {"_id": 0},
    ).to_list(10000)
    open_opportunities = await db.opportunities.find(
        {"status": {"$in": ["aberta", "em_analise"]}, "expected_close_date": {"$exists": True, "$ne": None}},
        {"_id": 0},
    ).to_list(10000)
    clients = {item["id"]: item.get("name", "") for item in await db.clients.find({}, {"_id": 0}).to_list(2000)}
    opportunities = {item["id"]: item.get("description", "") for item in await db.opportunities.find({}, {"_id": 0}).to_list(10000)}
    today = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
    anchor = today if year == today.year else datetime(year, 1, 1)
    thirty_days_end = anchor + timedelta(days=30)
    next_quarter_month = ((anchor.month - 1) // 3 + 1) * 3 + 1
    quarter_end = (datetime(year + 1, 1, 1) if next_quarter_month > 12 else datetime(year, next_quarter_month, 1)) - timedelta(days=1)
    year_end = datetime(year, 12, 31)
    proposal_labels = {"em_elaboracao": "Em elaboração", "enviada": "Enviada", "em_negociacao": "Em negociação"}
    opportunity_labels = {"aberta": "Aberta", "em_analise": "Em análise"}
    records = []

    def add_record(kind: str, record: dict, due_value: object, value: float, vab: float, probability: float, status: str, description: str):
        due_date = str(due_value or "")[:10]
        try:
            date = datetime.strptime(due_date, "%Y-%m-%d")
        except ValueError:
            return
        if date < anchor or date > year_end:
            return
        horizon = "Próximos 30 dias" if date <= min(thirty_days_end, year_end) else ("Trimestre atual" if date <= quarter_end else "Até final do ano")
        records.append({"horizon": horizon, "date": date, "type": kind, "number": record.get("number") or f"OPP-{str(record.get('id') or '-')[:8].upper()}", "client": clients.get(record.get("client_id"), "-"), "description": description, "value": value * probability / 100, "vab": vab * probability / 100, "probability": probability / 100, "status": status, "due_date": due_date})

    for proposal in proposals:
        probability = min(100, max(0, float(proposal.get("probability") if proposal.get("probability") is not None else 100)))
        add_record("Proposta", proposal, proposal.get("expected_close_date") or proposal.get("next_follow_up_date"), float(proposal.get("total_net") or 0), float(proposal.get("total_vab") or 0), probability, proposal_labels.get(proposal.get("status"), proposal.get("status", "-")), proposal.get("description") or opportunities.get(proposal.get("opportunity_id"), "-"))
    for opportunity in open_opportunities:
        probability = min(100, max(0, float(opportunity.get("probability") or 0)))
        add_record("Oportunidade", opportunity, opportunity.get("expected_close_date"), float(opportunity.get("estimated_value") or 0), float(opportunity.get("estimated_vab") or 0), probability, opportunity_labels.get(opportunity.get("status"), opportunity.get("status", "-")), opportunity.get("description", "-"))

    allowed = {"30d": {"Próximos 30 dias"}, "quarter": {"Próximos 30 dias", "Trimestre atual"}, "year": {"Próximos 30 dias", "Trimestre atual", "Até final do ano"}}[scope]
    records = sorted((record for record in records if record["horizon"] in allowed), key=lambda record: record["date"])
    return _xlsx_response(
        "Previsão de fecho",
        ["Horizonte", "Tipo", "Registo", "Cliente", "Descrição", "Valor s/ IVA (EUR)", "VAB (EUR)", "Probabilidade", "Estado", "Data prevista de fecho"],
        [[record["horizon"], record["type"], record["number"], record["client"], record["description"], record["value"], record["vab"], record["probability"], record["status"], record["due_date"]] for record in records],
        f"previsao-fecho-{year}-{scope}.xlsx",
        (20, 16, 20, 34, 54, 22, 20, 16, 20, 22),
        (6, 7),
        (8,),
    )


@router.get("/exports/dashboard.pdf")
async def export_dashboard_pdf(year: int = Query(datetime.now().year, ge=2000, le=2100), user: dict = Depends(get_current_user)):
    kpis = await _kpis_summary(year)
    fi = await _forecast_invoicing(year)
    fr = await _forecast_receiving(year)
    vab = await _vab_analysis(year)
    bc = (await by_commercial(user, year))["rows"]
    bcli = (await by_client(user, year))["rows"]
    bm = (await by_manufacturer(user, year))["rows"]
    # Filtrar "(sem fabricante)" para o snapshot
    bm = [r for r in bm if r.get("manufacturer_id")]
    return build_dashboard_pdf(kpis, fr, fi, vab, bc, bcli, bm)


@router.get("/exports/proposals-follow-up.pdf")
async def export_proposals_follow_up_pdf(year: int = Query(datetime.now().year, ge=2000, le=2100), scope: Optional[str] = Query(None, pattern="^(30d|quarter|year)$"), user: dict = Depends(get_current_user)):
    """Exporta a previsão de fecho em tabelas por horizonte."""
    proposals = await db.proposals.find(
        {"status": {"$in": ["em_elaboracao", "enviada", "em_negociacao"]}, "$or": [{"expected_close_date": {"$exists": True, "$ne": None}}, {"next_follow_up_date": {"$exists": True, "$ne": None}}]},
        {"_id": 0},
    ).to_list(10000)
    open_opportunities = await db.opportunities.find(
        {"status": {"$in": ["aberta", "em_analise"]}, "expected_close_date": {"$exists": True, "$ne": None}},
        {"_id": 0},
    ).to_list(10000)
    clients = {c["id"]: c.get("name", "") for c in await db.clients.find({}, {"_id": 0}).to_list(2000)}
    opportunities = {o["id"]: o.get("description", "") for o in await db.opportunities.find({}, {"_id": 0}).to_list(10000)}
    today = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
    anchor = today if year == today.year else datetime(year, 1, 1)
    thirty_days_end = anchor + timedelta(days=30)
    next_quarter_month = ((anchor.month - 1) // 3 + 1) * 3 + 1
    quarter_end = (datetime(year + 1, 1, 1) if next_quarter_month > 12 else datetime(year, next_quarter_month, 1)) - timedelta(days=1)
    year_end = datetime(year, 12, 31)
    status_labels = {
        "em_elaboracao": "Em elaboração", "enviada": "Enviada", "em_negociacao": "Em negociação",
    }
    opportunity_status_labels = {"aberta": "Aberta", "em_analise": "Em análise"}
    rows = []
    for proposal in proposals:
        follow_up = str(proposal.get("expected_close_date") or proposal.get("next_follow_up_date") or "")[:10]
        try:
            follow_up_date = datetime.strptime(follow_up, "%Y-%m-%d")
        except ValueError:
            continue
        if follow_up_date < anchor or follow_up_date > year_end:
            continue
        probability = min(100, max(0, float(proposal.get("probability") if proposal.get("probability") is not None else 100)))
        factor = probability / 100
        rows.append({
            "type": "Proposta",
            "number": proposal.get("number", "-"),
            "client": clients.get(proposal.get("client_id"), "-"),
            "description": opportunities.get(proposal.get("opportunity_id"), "-"),
            "value": float(proposal.get("total_net") or 0) * factor,
            "vab": float(proposal.get("total_vab") or 0) * factor,
            "probability": probability,
            "status": status_labels.get(proposal.get("status"), proposal.get("status", "-")),
            "follow_up": follow_up,
            "date": follow_up_date,
        })
    for opportunity in open_opportunities:
        follow_up = str(opportunity.get("expected_close_date") or "")[:10]
        try:
            follow_up_date = datetime.strptime(follow_up, "%Y-%m-%d")
        except ValueError:
            continue
        if follow_up_date < anchor or follow_up_date > year_end:
            continue
        probability = min(100, max(0, float(opportunity.get("probability") or 0)))
        factor = probability / 100
        rows.append({
            "type": "Oportunidade",
            "number": f"OPP-{str(opportunity.get('id') or '-')[:8].upper()}",
            "client": clients.get(opportunity.get("client_id"), "-"),
            "description": opportunity.get("description", "-"),
            "value": float(opportunity.get("estimated_value") or 0) * factor,
            "vab": float(opportunity.get("estimated_vab") or 0) * factor,
            "probability": probability,
            "status": opportunity_status_labels.get(opportunity.get("status"), opportunity.get("status", "-")),
            "follow_up": follow_up,
            "date": follow_up_date,
        })
    rows.sort(key=lambda item: item["date"])
    for row in rows:
        row["horizon"] = "Próximos 30 dias" if row["date"] <= min(thirty_days_end, year_end) else ("Trimestre atual" if row["date"] <= quarter_end else "Até final do ano")
    if scope:
        allowed = {"30d": {"Próximos 30 dias"}, "quarter": {"Próximos 30 dias", "Trimestre atual"}, "year": {"Próximos 30 dias", "Trimestre atual", "Até final do ano"}}[scope]
        title = {"30d": "PRÓXIMOS 30 DIAS", "quarter": "QUARTER", "year": "ATÉ AO FINAL DO ANO"}[scope]
        return build_proposal_follow_up_pdf(year, [(title, [row for row in rows if row["horizon"] in allowed])])
    sections = [
        ("PRÓXIMOS 30 DIAS", [row for row in rows if row["date"] <= min(thirty_days_end, year_end)]),
        ("QUARTER", [row for row in rows if row["date"] <= quarter_end]),
        ("ATÉ AO FINAL DO ANO", rows),
    ]
    return build_proposal_follow_up_pdf(year, sections)

