from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, ListFlowable, ListItem, PageBreak, HRFlowable


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "MANUAL_USO_WHYMOB.md"
OUT_DOCX = ROOT / "docs" / "MANUAL_USO_WHYMOB.docx"
OUT_PDF = ROOT / "docs" / "MANUAL_USO_WHYMOB.pdf"

BRAND_BLUE = RGBColor(0x00, 0x2F, 0xA7)
BRAND_GREEN = RGBColor(0x00, 0xA8, 0x59)
BRAND_RED = RGBColor(0xFF, 0x2A, 0x00)
PDF_BLUE = colors.HexColor("#002FA7")
PDF_GREEN = colors.HexColor("#00A859")
PDF_RED = colors.HexColor("#FF2A00")


def parse_markdown(text: str):
    blocks = []
    paragraph_buffer = []
    bullet_buffer = []

    def flush_paragraph():
        nonlocal paragraph_buffer
        if paragraph_buffer:
            blocks.append(("p", " ".join(paragraph_buffer).strip()))
            paragraph_buffer = []

    def flush_bullets():
        nonlocal bullet_buffer
        if bullet_buffer:
            blocks.append(("ul", bullet_buffer[:]))
            bullet_buffer = []

    for raw_line in text.splitlines():
        line = raw_line.rstrip()
        stripped = line.strip()

        if not stripped:
            flush_paragraph()
            flush_bullets()
            continue

        if stripped.startswith("# "):
            flush_paragraph()
            flush_bullets()
            blocks.append(("h1", stripped[2:].strip()))
            continue

        if stripped.startswith("## "):
            flush_paragraph()
            flush_bullets()
            blocks.append(("h2", stripped[3:].strip()))
            continue

        if stripped.startswith("### "):
            flush_paragraph()
            flush_bullets()
            blocks.append(("h3", stripped[4:].strip()))
            continue

        if stripped.startswith("- "):
            flush_paragraph()
            bullet_buffer.append(stripped[2:].strip())
            continue

        paragraph_buffer.append(stripped)

    flush_paragraph()
    flush_bullets()
    return blocks


def shade_cell(cell, fill_hex: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill_hex)
    tc_pr.append(shd)


def create_docx(blocks):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    cover = doc.add_table(rows=1, cols=2)
    cover.columns[0].width = Cm(2.5)
    cover.columns[1].width = Cm(13.5)
    shade_cell(cover.cell(0, 0), "002FA7")
    cover.cell(0, 0).text = ""
    p = cover.cell(0, 1).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("WhyMob CRM\n")
    r.font.size = Pt(24)
    r.font.bold = True
    r.font.color.rgb = BRAND_BLUE
    r2 = p.add_run("Manual de Uso da Ferramenta")
    r2.font.size = Pt(16)
    r2.font.bold = True

    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.add_run("Identidade visual de referência: ").bold = True
    meta.add_run("azul #002FA7, verde #00A859, tipografia limpa e orientada à operação.")
    meta.style = doc.styles["Normal"]

    doc.add_paragraph()

    for kind, content in blocks:
        if kind == "h1":
            para = doc.add_paragraph()
            run = para.add_run(content)
            run.font.size = Pt(20)
            run.font.bold = True
            run.font.color.rgb = BRAND_BLUE
            para.space_after = Pt(10)
        elif kind == "h2":
            para = doc.add_paragraph()
            run = para.add_run(content)
            run.font.size = Pt(15)
            run.font.bold = True
            run.font.color.rgb = BRAND_BLUE
            para.space_before = Pt(10)
            para.space_after = Pt(6)
        elif kind == "h3":
            para = doc.add_paragraph()
            run = para.add_run(content)
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = BRAND_GREEN
            para.space_before = Pt(8)
            para.space_after = Pt(4)
        elif kind == "p":
            para = doc.add_paragraph(content)
            para.style = doc.styles["Normal"]
            para.paragraph_format.space_after = Pt(5)
        elif kind == "ul":
            for item in content:
                para = doc.add_paragraph(style="List Bullet")
                para.add_run(item)

    doc.save(OUT_DOCX)


def create_pdf(blocks):
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "WhyTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=22,
        textColor=PDF_BLUE,
        spaceAfter=8,
        alignment=TA_CENTER,
    )
    subtitle_style = ParagraphStyle(
        "WhySubtitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=13,
        textColor=PDF_GREEN,
        spaceAfter=16,
        alignment=TA_CENTER,
    )
    h1_style = ParagraphStyle(
        "WhyH1",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=18,
        textColor=PDF_BLUE,
        spaceBefore=12,
        spaceAfter=8,
    )
    h2_style = ParagraphStyle(
        "WhyH2",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=14,
        textColor=PDF_BLUE,
        spaceBefore=10,
        spaceAfter=6,
    )
    h3_style = ParagraphStyle(
        "WhyH3",
        parent=styles["Heading3"],
        fontName="Helvetica-Bold",
        fontSize=11.5,
        textColor=PDF_GREEN,
        spaceBefore=8,
        spaceAfter=4,
    )
    body_style = ParagraphStyle(
        "WhyBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=10.5,
        leading=15,
        textColor=colors.HexColor("#111111"),
        spaceAfter=6,
    )

    story = [
        Spacer(1, 1.8 * cm),
        Paragraph("WhyMob CRM", title_style),
        Paragraph("Manual de Uso da Ferramenta", subtitle_style),
        HRFlowable(width="100%", thickness=2, color=PDF_BLUE),
        Spacer(1, 0.6 * cm),
        Paragraph("Documento de utilização funcional com a identidade visual base da aplicação.", body_style),
        PageBreak(),
    ]

    for kind, content in blocks:
        if kind == "h1":
            story.append(Paragraph(content, h1_style))
        elif kind == "h2":
            story.append(Paragraph(content, h2_style))
        elif kind == "h3":
            story.append(Paragraph(content, h3_style))
        elif kind == "p":
            story.append(Paragraph(content, body_style))
        elif kind == "ul":
            items = [ListItem(Paragraph(item, body_style)) for item in content]
            story.append(ListFlowable(items, bulletType="bullet", leftIndent=14))
            story.append(Spacer(1, 0.15 * cm))

    doc = SimpleDocTemplate(
        str(OUT_PDF),
        pagesize=A4,
        rightMargin=2 * cm,
        leftMargin=2 * cm,
        topMargin=1.8 * cm,
        bottomMargin=1.8 * cm,
    )
    doc.build(story)


def main():
    text = SOURCE.read_text(encoding="utf-8")
    blocks = parse_markdown(text)
    create_docx(blocks)
    create_pdf(blocks)
    print(f"DOCX: {OUT_DOCX}")
    print(f"PDF: {OUT_PDF}")


if __name__ == "__main__":
    main()
